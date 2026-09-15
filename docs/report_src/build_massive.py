import os
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def setup_styles(doc):
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    
    # Heading 1
    h1 = doc.styles['Heading 1']
    h1.font.name = 'Times New Roman'
    h1.font.size = Pt(18)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    
    # Heading 2
    h2 = doc.styles['Heading 2']
    h2.font.name = 'Times New Roman'
    h2.font.size = Pt(16)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    
    # Heading 3
    h3 = doc.styles['Heading 3']
    h3.font.name = 'Times New Roman'
    h3.font.size = Pt(14)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0, 0, 0)

def add_cover(doc):
    for _ in range(5): doc.add_paragraph()
    p = doc.add_heading("BÁO CÁO ĐỒ ÁN NGÀNH", level=1)
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    for _ in range(2): doc.add_paragraph()
    p = doc.add_heading("AI LEGAL INTELLIGENCE PLATFORM", level=2)
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    for _ in range(10): doc.add_paragraph()
    p = doc.add_paragraph("Sinh viên thực hiện: Vũ Lưu Minh Toán - 2354050139")
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p = doc.add_paragraph("Giảng viên hướng dẫn: Hồ Hướng Thiện")
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_page_break()

def add_table(doc, lines):
    if not lines: return
    # Find max columns
    table_data = [line.split('|') for line in lines]
    cols = max(len(row) for row in table_data)
    if cols <= 1:
        # Not a real table, just text with |
        for line in lines:
            doc.add_paragraph(line)
        return
        
    table = doc.add_table(rows=len(table_data), cols=cols)
    table.style = 'Table Grid'
    
    for r_idx, row in enumerate(table_data):
        for c_idx, cell in enumerate(row):
            if c_idx < cols:
                p = table.cell(r_idx, c_idx).paragraphs[0]
                run = p.add_run(cell.strip())
                run.font.name = 'Times New Roman'
                run.font.size = Pt(12)
                if r_idx == 0:
                    run.font.bold = True
                p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER if r_idx == 0 else WD_PARAGRAPH_ALIGNMENT.LEFT

def parse_and_add_text(doc, text):
    lines = text.split('\n')
    current_table_lines = []
    in_table = False
    
    for line in lines:
        line = line.strip()
        if not line:
            if in_table:
                add_table(doc, current_table_lines)
                current_table_lines = []
                in_table = False
            continue
            
        if line.startswith("--- Tables:") or line.startswith("==="):
            if in_table:
                add_table(doc, current_table_lines)
                current_table_lines = []
                in_table = False
            continue
            
        if line.startswith("Table "):
            if in_table:
                add_table(doc, current_table_lines)
                current_table_lines = []
                in_table = False
            doc.add_heading(line.replace(":", ""), level=3)
            continue
            
        if '|' in line:
            in_table = True
            current_table_lines.append(line)
        else:
            if in_table:
                add_table(doc, current_table_lines)
                current_table_lines = []
                in_table = False
                
            if re.match(r'^(CHƯƠNG|PHẦN)\s+\d+', line, re.IGNORECASE):
                doc.add_heading(line, level=1)
            elif re.match(r'^\d+\.\s+[A-ZÀÁÂÃÈÉÊÌÍÒÓÔÕÙÚĂĐĨŨƠƯĂẠẢẤẦẨẪẬẮẰẲẴẶẸẺẼỀỀỂỄỆỈỊỌỎỐỒỔỖỘỚỜỞỠỢỤỦỨỪỬỮỰỲỴÝỶỸ]', line):
                doc.add_heading(line, level=2)
            elif re.match(r'^\d+\.\d+\.', line):
                doc.add_heading(line, level=3)
            else:
                p = doc.add_paragraph(line)
                p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
                
    if in_table and current_table_lines:
        add_table(doc, current_table_lines)

def add_code_appendix(doc, base_dir, folder_path):
    full_path = os.path.join(base_dir, folder_path)
    if not os.path.exists(full_path):
        return
        
    for root, dirs, files in os.walk(full_path):
        for file in files:
            if file.endswith('.py') or file.endswith('.ts') or file.endswith('.tsx'):
                file_path = os.path.join(root, file)
                rel_path = os.path.relpath(file_path, start=base_dir)
                doc.add_heading(f"File: {rel_path}", level=3)
                try:
                    with open(file_path, 'r', encoding='utf-8') as f:
                        code = f.read()
                    p = doc.add_paragraph(code)
                    for run in p.runs:
                        run.font.name = 'Courier New'
                        run.font.size = Pt(10)
                except Exception:
                    pass

def main():
    print("Start building massive report...")
    doc = Document()
    setup_styles(doc)
    add_cover(doc)
    
    docs_file = r"D:\NCKH\docs_extracted.txt"
    if os.path.exists(docs_file):
        print("Injecting analysis & design documents...")
        with open(docs_file, 'r', encoding='utf-8') as f:
            lines = f.readlines()
            
        fixed_lines = []
        for line in lines:
            try:
                # Nếu line bị mojibake từ CP437
                fixed_line = line.encode('cp437').decode('utf-8')
                fixed_lines.append(fixed_line)
            except Exception:
                try:
                    # Nếu line bị mojibake từ CP1252
                    fixed_line = line.encode('cp1252').decode('utf-8')
                    fixed_lines.append(fixed_line)
                except Exception:
                    # Nếu line đã chuẩn hoặc không thể fix, giữ nguyên
                    fixed_lines.append(line)
        text = ''.join(fixed_lines)
                
        doc.add_heading("PHẦN 1: NỘI DUNG PHÂN TÍCH VÀ THIẾT KẾ", level=1)
        parse_and_add_text(doc, text)
    else:
        print(f"Error: Could not find {docs_file}")
        
    status_file = r"D:\NCKH\PROJECT_STATUS.md"
    if os.path.exists(status_file):
        print("Injecting project status...")
        doc.add_page_break()
        doc.add_heading("PHẦN 2: TIẾN ĐỘ VÀ TRẠNG THÁI DỰ ÁN", level=1)
        with open(status_file, 'r', encoding='utf-8') as f:
            text = f.read()
        parse_and_add_text(doc, text)

    # Append Source Code Appendix
    print("Injecting source code appendices...")
    doc.add_page_break()
    doc.add_heading("PHỤ LỤC: CHI TIẾT MÃ NGUỒN (SOURCE CODE)", level=1)
    
    base_backend = r"D:\NCKH\backend"
    doc.add_heading("Phụ lục 1: Cấu trúc Core và Router (FastAPI)", level=2)
    add_code_appendix(doc, base_backend, r"app\core")
    add_code_appendix(doc, base_backend, r"app\routers")
    
    doc.add_heading("Phụ lục 2: Cấu trúc Models và Schemas (Database)", level=2)
    add_code_appendix(doc, base_backend, r"app\models")
    add_code_appendix(doc, base_backend, r"app\schemas")
    
    doc.add_heading("Phụ lục 3: Cấu trúc Services (Xử lý Logic RAG & Vector)", level=2)
    add_code_appendix(doc, base_backend, r"app\services")
    
    # Save the huge doc
    output_path = r"D:\NCKH\docs\Bao_cao_Do_An_Nganh_MASSIVE.docx"
    doc.save(output_path)
    print(f"Successfully saved to: {output_path} (Open this file for 80-100 pages!)")

if __name__ == "__main__":
    main()
