from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def setup_styles(doc):
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    
    # Heading 1
    h1 = doc.styles['Heading 1']
    h1.font.name = 'Times New Roman'
    h1.font.size = Pt(18)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    
    # Heading 2
    h2 = doc.styles['Heading 2']
    h2.font.name = 'Times New Roman'
    h2.font.size = Pt(16)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    
    # Heading 3
    h3 = doc.styles['Heading 3']
    h3.font.name = 'Times New Roman'
    h3.font.size = Pt(14)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0, 0, 0)

def add_heading(doc, text, level=1, align=None):
    p = doc.add_heading(text, level=level)
    if align is not None:
        p.alignment = align
    return p

def add_paragraph(doc, text, bold=False, italic=False, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    return p

def add_bullet(doc, text):
    doc.add_paragraph(text, style='List Bullet')

def add_numbered(doc, text):
    doc.add_paragraph(text, style='List Number')

def add_code(doc, text):
    p = doc.add_paragraph(style='Normal')
    run = p.add_run(text)
    run.font.name = 'Courier New'
    run.font.size = Pt(11)
    
def add_page_break(doc):
    doc.add_page_break()

def add_cover(doc):
    add_paragraph(doc, "\n\n\n\n\n", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_heading(doc, "BÁO CÁO ĐỒ ÁN NGÀNH", level=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_heading(doc, "TÊN ĐỀ TÀI:", level=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_heading(doc, "AI LEGAL INTELLIGENCE PLATFORM (AILIP)\nNền tảng Tra cứu và Phân tích Pháp luật Việt Nam Ứng dụng AI", level=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    add_paragraph(doc, "\n\n\n\n\n", align=WD_ALIGN_PARAGRAPH.CENTER)
    p = add_paragraph(doc, "Sinh viên thực hiện: ...", align=WD_ALIGN_PARAGRAPH.CENTER)
    p = add_paragraph(doc, "Giảng viên hướng dẫn: ...", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_page_break(doc)
