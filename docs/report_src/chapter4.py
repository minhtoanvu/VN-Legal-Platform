from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def add_chapter_4(doc, format_heading, format_normal):
    """
    Chương 4: HIỆN THỰC HỆ THỐNG - API, UI/UX, Code Snippets và CI/CD
    """
    doc.add_page_break()
    h = doc.add_paragraph('Chương 4. HIỆN THỰC VÀ TRIỂN KHAI HỆ THỐNG')
    format_heading(h, level=1)

    # 4.1 API
    h = doc.add_paragraph('4.1. Đặc tả Giao diện Lập trình (RESTful APIs)')
    format_heading(h, level=2)
    p = doc.add_paragraph(
        "Hệ thống Backend được phát triển bằng FastAPI, tuân thủ nghiêm ngặt tiêu chuẩn RESTful API. "
        "Mọi endpoint đều được đặt dưới tiền tố '/api/v1/' để đảm bảo tính versioning. Dưới đây là các API cốt lõi:"
    )
    format_normal(p)

    apis = [
        ("POST /api/v1/auth/login", "Xác thực người dùng, trả về JWT Access Token (30 phút) và Refresh Token."),
        ("GET /api/v1/search?q={query}&mode={hybrid|semantic|keyword}", "Thực thi truy vấn văn bản, trả về danh sách tài liệu kèm RRF score."),
        ("POST /api/v1/ai/chat", "Endpoint Server-Sent Events (SSE) nhận câu hỏi tự nhiên và stream câu trả lời RAG thời gian thực."),
        ("GET /api/v1/graph/{doc_id}?depth=2", "Truy xuất danh sách Nodes và Edges của Đồ thị Tri thức để Vis.js render."),
        ("GET /api/v1/analytics/dashboard", "Thực thi Aggregation queries trên dữ liệu PostgreSQL, trả về dữ liệu vẽ 5 biểu đồ Recharts.")
    ]
    
    for api, desc in apis:
        p = doc.add_paragraph(f"• {api}: {desc}", style='List Bullet')
        format_normal(p)

    p = doc.add_paragraph(
        "\nCấu trúc phản hồi lỗi (Error Response) được chuẩn hóa dưới dạng JSON thống nhất toàn hệ thống:\n"
        '{"error": {"code": "RATE_LIMIT_EXCEEDED", "message": "Bạn đã vượt quá 15 câu hỏi/phút."}}'
    )
    format_normal(p)

    # 4.2 Code Snippets
    h = doc.add_paragraph('4.2. Hiện thực Thuật toán Lõi (Code Snippets)')
    format_heading(h, level=2)
    p = doc.add_paragraph(
        "Để chứng minh quá trình hiện thực, dưới đây là mã nguồn Python xử lý thuật toán Reciprocal Rank Fusion (RRF) "
        "nhằm kết hợp kết quả của Semantic Search và BM25:"
    )
    format_normal(p)

    p_code = doc.add_paragraph(
        "def rrf_score(rank: int, k: int = 60) -> float:\n"
        "    # Tính điểm RRF dựa trên hạng (rank) thay vì điểm tuyệt đối\n"
        "    return 1.0 / (k + rank)\n\n"
        "def hybrid_rerank(bm25_results: list, semantic_results: list) -> list:\n"
        "    scores = {}\n"
        "    for rank, doc_id in enumerate(bm25_results):\n"
        "        scores[doc_id] = scores.get(doc_id, 0) + rrf_score(rank)\n"
        "    for rank, doc_id in enumerate(semantic_results):\n"
        "        scores[doc_id] = scores.get(doc_id, 0) + rrf_score(rank)\n"
        "    # Sắp xếp lại theo tổng điểm RRF giảm dần\n"
        "    return sorted(scores.keys(), key=lambda x: scores[x], reverse=True)\n"
    )
    for run in p_code.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(11)

    # 4.3 UI/UX Design System
    h = doc.add_paragraph('4.3. Thiết kế Giao diện (UI/UX Design System)')
    format_heading(h, level=2)
    p = doc.add_paragraph(
        "Trái ngược với giao diện bảng biểu khô khan của các website pháp luật hiện hành, AILIP áp dụng phong cách "
        "thiết kế 'Dark Glassmorphism' (Kính mờ trên nền tối). Phong cách này mang lại sự chuyên nghiệp, công nghệ cao "
        "và rất phù hợp cho doanh nghiệp.\n\n"
        "- Màu nền: Navy/Charcoal tối để giảm mỏi mắt khi đọc văn bản luật dài.\n"
        "- Hiệu ứng: Backdrop blur (làm mờ nền) và viền phát sáng (neon glow) cho các thành phần tương tác như Khung Chat AI "
        "hay Card thống kê Dashboard.\n"
        "- Typography: Sử dụng font Sans-serif hiện đại, độ tương phản cao, tối ưu khả năng đọc (Readability).\n"
        "- Trạng thái loading: Hệ thống áp dụng Skeleton UI thay vì Spinner truyền thống, giúp người dùng cảm nhận tốc độ "
        "phản hồi nhanh hơn trong lúc chờ AI sinh câu trả lời."
    )
    format_normal(p)

    # 4.4 DevOps & CI/CD
    h = doc.add_paragraph('4.4. Môi trường Triển khai và CI/CD')
    format_heading(h, level=2)
    p = doc.add_paragraph(
        "Dự án tuân thủ quy trình phát triển hiện đại (QAOps):\n"
        "1. Dockerization: Toàn bộ Backend, Frontend và PostgreSQL+pgvector được đóng gói trong 'docker-compose.yml'. "
        "Khống chế bộ nhớ RAM của database bằng 'deploy.resources.limits.memory' để tránh sập server.\n"
        "2. Continuous Integration (CI): Cấu hình GitHub Actions ('backend-ci.yml', 'e2e-ci.yml') tự động chạy:\n"
        "   - Unit Tests (pytest) cho Backend API.\n"
        "   - Linter (ESLint) và Type-check cho Frontend TypeScript.\n"
        "   - End-to-End Tests (Playwright) trên Node.js v20 kiểm tra toàn bộ luồng đăng nhập và chat RAG.\n"
        "Mọi commit lên nhánh chính đều phải pass 100% các bài test này trước khi được merge."
    )
    format_normal(p)
